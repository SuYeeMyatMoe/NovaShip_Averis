"""Excel export, 3-action auto-draft, operator warnings, mocked Gemini LLM/OCR."""
from __future__ import annotations

import base64
import io
import os
import sys
from types import ModuleType, SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import load_workbook

os.environ.setdefault("REPO_BACKEND", "memory")
os.environ["AUTO_SEED"] = "0"
os.environ["LLM_PROVIDER"] = "none"

from app.ai.operator_behaviour import clear_operator_burst  # noqa: E402
from app.main import app  # noqa: E402
from app.readers.document_reader import read_document  # noqa: E402

client = TestClient(app)
SUP = {"X-User-Id": "u_sup_1"}
OPS = {"X-User-Id": "u_ops_1"}

SI_TXT = """SHIPPING INSTRUCTION
Shipper/Exporter: APRIL FAR EAST (M) SDN BHD
CONSIGNEE: MOORIM SP CO., LTD
NOTIFY PARTY: UAB NOVAKOPA
Port of Loading: PORT KLANG (WESTPORT), MALAYSIA (MYPKG)
Discharge Port: CALLAO, PERU (PECLL)
No. of Containers or Packages: 3 x 40'HC
Gross Weight (KG): 22,000 KG
"""
BL_TXT = """BILL OF LADING (DRAFT)
SHIPPER: APRIL FAR EAST (M) SDN BHD
To the Order of: MOORIM SP CO., LTD
Notify: UAB NOVAKOPA
Load Port: PORT KLANG (WESTPORT), MALAYSIA (MYPKG)
POD: CALLAO, PERU (PECLL)
Container Count: 4 x 40'HC
Gross Wt (kgs): 22,000 KG
"""


def _ingest(email_id: str) -> dict:
    payload = {
        "email_id": email_id,
        "from": "docs@vitalsolutions.sg",
        "subject": f"TO CONFIRM DOCS _ {email_id}",
        "body": "Attached are the SI and draft BL. Please confirm.",
        "provider": "test",
        "attachments": [
            {"name": f"{email_id}_SI.txt", "content_base64": base64.b64encode(SI_TXT.encode()).decode()},
            {"name": f"{email_id}_BL.txt", "content_base64": base64.b64encode(BL_TXT.encode()).decode()},
        ],
    }
    r = client.post("/webhooks/email", json=payload, headers=SUP)
    assert r.status_code == 201, r.text
    return r.json()


def test_export_xlsx_contains_expected_sheets_and_selected_ids():
    row = _ingest("xlsx_001")
    cid = row["id"]
    r = client.get("/export/cases.xlsx", headers=SUP)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Cases", "Field results", "Summary"]
    assert [c.value for c in wb["Cases"][1]][0] == "case_id"
    case_ids = [cell.value for cell in wb["Cases"]["A"][1:] if cell.value]
    assert cid in case_ids
    assert wb["Summary"]["A2"].value == "cases"
    batch = client.post("/cases/batch", json={"action": "export_xlsx", "case_ids": [cid]}, headers=SUP).json()
    assert batch["filename"] == "cases.xlsx"
    selected = load_workbook(io.BytesIO(base64.b64decode(batch["xlsx_base64"])))
    ids = [cell.value for cell in selected["Cases"]["A"][1:] if cell.value]
    assert ids == [cid]
    field_ids = {cell.value for cell in selected["Field results"]["A"][1:] if cell.value}
    assert cid in field_ids
    assert client.get("/export/cases.xlsx", headers=OPS).status_code == 403


def test_third_operator_action_saves_draft_and_never_sends():
    clear_operator_burst()
    cid = _ingest("autodraft_001")["id"]
    seeded = client.get(f"/cases/{cid}", headers=SUP).json()
    assert seeded["drafts"]
    rejected = client.post(f"/cases/{cid}/reject", json={"draft_id": seeded["drafts"][0]["id"], "note": "rework"}, headers=SUP)
    assert rejected.status_code == 200
    assert not any(a["signal"] == "AUTO_DRAFT_AFTER_REPEATED_ACTIONS" for a in rejected.json()["anomalies"])
    for _ in range(2):
        r = client.post(f"/cases/{cid}/assign", json={"user_id": "u_ops_1"}, headers=SUP)
        assert r.status_code == 200
        last = r
    assert any(a["signal"] == "AUTO_DRAFT_AFTER_REPEATED_ACTIONS" for a in last.json()["anomalies"])
    body = last.json()
    assert body["drafts"]
    assert all(d["status"] not in {"SENT", "SIMULATED"} for d in body["drafts"])
    audit = client.get(f"/cases/{cid}/audit", headers=SUP).json()
    assert any(e["action"] == "AUTO_DRAFT_AFTER_REPEATED_ACTIONS" for e in audit["events"])
    queue = client.get("/security/queue", headers=SUP).json()
    assert any(item["case_id"] == cid and any(a["signal"] == "AUTO_DRAFT_AFTER_REPEATED_ACTIONS" for a in item["anomalies"]) for item in queue["items"])


def test_live_draft_skips_auto_draft_on_third_action():
    clear_operator_burst()
    cid = _ingest("skipdraft_001")["id"]
    drafted = client.post(f"/cases/{cid}/draft", json={}, headers=SUP)
    assert drafted.status_code == 200
    n0 = len(drafted.json()["drafts"])
    for _ in range(3):
        assert client.post(f"/cases/{cid}/assign", json={"user_id": "u_ops_1"}, headers=SUP).status_code == 200
    later = client.get(f"/cases/{cid}", headers=SUP).json()
    assert len(later["drafts"]) == n0
    assert not any(a["signal"] == "AUTO_DRAFT_AFTER_REPEATED_ACTIONS" for a in later["anomalies"])


def test_share_denied_records_operator_warning_and_stays_403():
    cid = _ingest("sharewarn_001")["id"]
    r = client.post(
        f"/cases/{cid}/share",
        json={"recipient_type": "NOTIFY_PARTY_CONTACT", "recipient_party_id": "p_safqa", "confirm_external": True},
        headers=OPS,
    )
    assert r.status_code == 403
    assert r.json()["detail"]["operator_warning"] is True
    case = client.get(f"/cases/{cid}", headers=SUP).json()
    assert any(a["signal"] == "OPERATOR_SHARE_DENIED" for a in case["anomalies"])
    audit = client.get(f"/cases/{cid}/audit", headers=SUP).json()
    assert any(e["action"] == "UNUSUAL_OPERATOR_BEHAVIOUR" for e in audit["events"])


def test_operator_burst_and_rapid_archive_warnings():
    clear_operator_burst()
    cid = _ingest("burst_001")["id"]
    last = None
    for _ in range(10):
        last = client.post(f"/cases/{cid}/assign", json={"user_id": "u_ops_1"}, headers=SUP)
        assert last.status_code == 200
    assert any(a["signal"] == "OPERATOR_BURST_MUTATIONS" for a in last.json()["anomalies"])

    ids = []
    for i in range(5):
        payload = {
            "email_id": f"archive_{i:03d}",
            "from": "rpa.bot@aprilasia.com",
            "subject": f"_RPA_ archive batch {i}",
            "body": "Automated notification. No action required.\n\n-- RPA Bot",
            "attachments": [],
        }
        row = client.post("/webhooks/email", json=payload, headers=SUP)
        assert row.status_code == 201, row.text
        ids.append(row.json()["id"])
    batch = client.post("/cases/batch", json={"action": "archive", "case_ids": ids}, headers=SUP)
    assert batch.status_code == 200
    first = client.get(f"/cases/{ids[0]}", headers=SUP).json()
    assert any(a["signal"] == "OPERATOR_RAPID_ARCHIVE" for a in first["anomalies"])


def test_repeated_login_failures_warn_without_locking():
    email = "burst.login@aprilasia.com"
    for i in range(3):
        r = client.post("/auth/login", json={"email": email, "password": "wrong-password"})
        assert r.status_code == 401
        if i == 2:
            assert r.json()["detail"].get("operator_warning") is True
    audit = client.get("/audit", headers=SUP).json()
    assert any(e["action"] == "UNUSUAL_OPERATOR_BEHAVIOUR" and (e.get("after") or {}).get("signal") == "OPERATOR_REPEATED_LOGIN_FAILURE" for e in audit["events"])
    ok = client.post("/auth/login", json={"email": "hari_mardianto@aprilasia.com", "password": os.environ.get("DEMO_PASSWORD", "novaship123")})
    assert ok.status_code == 200


def test_gemini_llm_mocked_and_falls_back_without_key(monkeypatch):
    import app.ai.llm as llm

    monkeypatch.setenv("LLM_PROVIDER", "gemini")
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    llm.reset_llm()
    disabled = llm.LLMClient()
    assert not disabled.enabled and disabled.complete("sys", "user") is None

    class FakeGemini:
        def __init__(self, **kwargs):
            self.kwargs = kwargs

        def invoke(self, messages):
            return SimpleNamespace(content='{"answer":"ok"}', usage_metadata={"input_tokens": 2, "output_tokens": 3})

    monkeypatch.setenv("GOOGLE_API_KEY", "not-a-real-key")
    monkeypatch.setattr("langchain_google_genai.ChatGoogleGenerativeAI", FakeGemini)
    llm.reset_llm()
    enabled = llm.LLMClient()
    assert enabled.enabled and enabled.provider == "gemini"
    result = enabled.complete("system", "user question")
    assert result is not None and result.provider == "gemini" and "ok" in result.text
    parsed = enabled.complete_json("system", "user")
    assert parsed == {"answer": "ok"}


def test_ocr_prefers_gemini_then_tesseract(monkeypatch):
    monkeypatch.setenv("OCR_ENABLED", "1")
    monkeypatch.setenv("GOOGLE_API_KEY", "not-a-real-key")
    monkeypatch.setattr("app.readers.document_reader._ocr_pdf_gemini", lambda data: "[[PAGE 1]]\nOCR FROM GEMINI")
    from app.readers import document_reader as dr

    text = dr._ocr_pdf(b"%PDF-1.4 fake")
    assert "OCR FROM GEMINI" in text

    monkeypatch.setattr("app.readers.document_reader._ocr_pdf_gemini", lambda data: "")
    fake_pdf = ModuleType("pdf2image")
    fake_pdf.convert_from_bytes = lambda data: [object()]
    fake_tess = ModuleType("pytesseract")
    fake_tess.image_to_string = lambda im: "TESSERACT TEXT"
    monkeypatch.setitem(sys.modules, "pdf2image", fake_pdf)
    monkeypatch.setitem(sys.modules, "pytesseract", fake_tess)
    assert "TESSERACT TEXT" in dr._ocr_pdf(b"%PDF-1.4 fake")

    monkeypatch.setattr("app.readers.document_reader._ocr_pdf", lambda data: "[[PAGE 1]]\nOCR FROM GEMINI")

    class _Page:
        def extract_text(self):
            return ""

    class _Reader:
        def __init__(self, *args, **kwargs):
            self.pages = [_Page()]

    monkeypatch.setattr("pypdf.PdfReader", _Reader)
    result = read_document("scan.pdf", b"%PDF-1.5 empty-text-layer")
    assert result.status.value == "EXTRACTED"
    assert "lower confidence" in (result.note or "")
    assert "OCR FROM GEMINI" in result.text
